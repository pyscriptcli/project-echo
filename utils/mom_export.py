"""Document export utilities for Minutes of the Meeting (Word and PDF templates)."""
import os
import datetime
from io import BytesIO

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor
import pandas as pd
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, ListFlowable, ListItem
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

import components.doc_fonts as docfonts
docfonts.ensure_pdf_fonts()


def set_cell_shading(cell, color_hex: str):
    """Set background shading color for a python-docx table cell."""
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


# -------------------------------------------------------------
# TEMPLATE 1: Standard Corporate (Combined Table)
# -------------------------------------------------------------
def export_to_word_template_1(df: pd.DataFrame, meeting_details: dict, other_discussions: str) -> BytesIO:
    """Generate Word (.docx) for Template 1 (Standard Corporate Combined Table)."""
    template_files = ["MOM_Template.docx", "MOM Template.docx"]
    template_path = next((f for f in template_files if os.path.exists(f)), None)
    doc = Document(template_path) if template_path else Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("MINUTES OF THE MEETING")
    r_title.bold = True
    r_title.underline = True
    r_title.font.name = docfonts.WORD_BODY
    r_title.font.size = Pt(11)

    company_target = meeting_details.get("external_attendees", [])
    primary_client_rep = company_target[0] if company_target else meeting_details.get("company_name", "").strip() or "CLIENT"
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run(f"PRIME PHILIPPINES & {primary_client_rep.upper()}")
    r_sub.bold = True
    r_sub.font.name = docfonts.WORD_BODY
    r_sub.font.size = Pt(11)

    date_str = meeting_details.get("date", "____________")
    time_str = meeting_details.get("time_range", "")
    full_date = f"Date: {date_str}" + (f", {time_str}" if time_str.strip() else "")
    
    p_date = doc.add_paragraph(full_date)
    p_date.paragraph_format.space_after = Pt(2)
    for r in p_date.runs:
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(10)

    p_loc = doc.add_paragraph(f"Location: {meeting_details.get('location', '____________')}")
    p_loc.paragraph_format.space_after = Pt(2)
    for r in p_loc.runs:
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(10)

    prime_atts = meeting_details.get("prime_attendees", [])
    ext_atts = meeting_details.get("external_attendees", [])
    
    p_att = doc.add_paragraph()
    p_att.paragraph_format.space_after = Pt(2)
    p_att.paragraph_format.tab_stops.add_tab_stop(Inches(1.35), WD_TAB_ALIGNMENT.LEFT)
    r_att_label = p_att.add_run("Attended by:")
    r_att_label.font.name, r_att_label.font.size = docfonts.WORD_BODY, Pt(10)
    
    first_attendee = True
    for att in ext_atts:
        if not att.strip():
            continue
        p = p_att if first_attendee else doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if not first_attendee:
            p.paragraph_format.left_indent = Inches(1.35)
        else:
            p.add_run("\t")
        comp_label = f", {meeting_details.get('company_name')}" if meeting_details.get('company_name') else ""
        r = p.add_run(f"{att}{comp_label}")
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(10)
        first_attendee = False

    for att in prime_atts:
        p = p_att if first_attendee else doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if not first_attendee:
            p.paragraph_format.left_indent = Inches(1.35)
        else:
            p.add_run("\t")
        r = p.add_run(f"{att} – PRIME Philippines")
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(10)
        first_attendee = False

    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(4)
    p_line.paragraph_format.space_after = Pt(6)
    r_line = p_line.add_run("_________________________________________________________________________________")
    r_line.font.name, r_line.font.color.rgb = docfonts.WORD_BODY, RGBColor(105, 114, 125)

    client_display = meeting_details.get('company_name', '').strip() or "the Client"
    p_intro = doc.add_paragraph(f"During the meeting held last {date_str}, PRIME Philippines, represented by the attendee/s shown above, met with {client_display} to discuss opportunities for collaboration.")
    p_intro.paragraph_format.space_after = Pt(10)
    for r in p_intro.runs:
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(9.5)

    table = doc.add_table(rows=len(df)+1, cols=4)
    table.alignment, table.style, table.autofit, table.allow_autofit = WD_TABLE_ALIGNMENT.CENTER, "Table Grid", False, False
    col_widths = [Inches(2.5), Inches(2.2), Inches(1.1), Inches(1.2)]
    headers = ["Discussion Points", "Action Plan", "Indicative Delivery Date", "Person-in-charge"]
    
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width, cell.text = col_widths[i], header
        set_cell_shading(cell, "C9AB4C")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].font.bold, p.runs[0].font.size, p.runs[0].font.name = True, Pt(9), docfonts.WORD_BODY

    for i, row in df.iterrows():
        cells = table.rows[i+1].cells
        cells[0].text = f"{i+1}. {str(row.get('Discussion Points', ''))}"
        cells[1].text = str(row.get("Action Plan", ""))
        cells[2].text = str(row.get("Indicative Delivery Date", ""))
        cells[3].text = str(row.get("Person-in-charge", ""))
        for c_idx, cell in enumerate(cells):
            cell.width = col_widths[c_idx]
            p = cell.paragraphs[0]
            if c_idx in [2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if p.runs:
                p.runs[0].font.size, p.runs[0].font.name = Pt(8.5), docfonts.WORD_BODY

    doc.add_paragraph()
    p_note = doc.add_paragraph("*Note: The indicative delivery date serves as reference point and still subject to changes. Furthermore, it depends on the progress of both parties.")
    p_note.paragraph_format.space_after = Pt(8)
    if p_note.runs:
        p_note.runs[0].italic, p_note.runs[0].font.name, p_note.runs[0].font.size = True, docfonts.WORD_BODY, Pt(8)

    if other_discussions.strip():
        p_od_head = doc.add_paragraph()
        p_od_head.paragraph_format.space_before, p_od_head.paragraph_format.space_after = Pt(6), Pt(4)
        r_od_head = p_od_head.add_run("Other Discussions:")
        r_od_head.bold, r_od_head.font.size, r_od_head.font.name = True, Pt(10), docfonts.WORD_BODY
        p_od = doc.add_paragraph(other_discussions)
        p_od.paragraph_format.space_after = Pt(12)
        for r in p_od.runs:
            r.font.name, r.font.size = docfonts.WORD_BODY, Pt(9.5)

    p_prep_label = doc.add_paragraph("Prepared by:")
    p_prep_label.paragraph_format.space_before = Pt(12)
    p_prep_label.paragraph_format.space_after = Pt(2)
    p_prep_label.runs[0].font.name, p_prep_label.runs[0].font.bold, p_prep_label.runs[0].font.size = docfonts.WORD_BODY, True, Pt(9.5)
    p_prep_line = doc.add_paragraph("_______________________________")
    p_prep_line.paragraph_format.space_after = Pt(2)
    p_prep_line.runs[0].font.name = docfonts.WORD_BODY
    prep_name = meeting_details.get("prep_name", "").strip() or "____________________"
    prep_desig = meeting_details.get("prep_desig", "").strip() or "PRIME Philippines"
    p_prep_info = doc.add_paragraph(f"{prep_name}\n{prep_desig}")
    p_prep_info.paragraph_format.space_after = Pt(12)
    for r in p_prep_info.runs:
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(9.5)

    p_conf_label = doc.add_paragraph("Confirmed by:")
    p_conf_label.paragraph_format.space_after = Pt(2)
    p_conf_label.runs[0].font.name, p_conf_label.runs[0].font.bold, p_conf_label.runs[0].font.size = docfonts.WORD_BODY, True, Pt(9.5)
    p_conf_line = doc.add_paragraph("_______________________________")
    p_conf_line.paragraph_format.space_after = Pt(2)
    p_conf_line.runs[0].font.name = docfonts.WORD_BODY
    conf_name = meeting_details.get("conf_name", "").strip() or (ext_atts[0] if ext_atts else "____________________")
    conf_desig = meeting_details.get("conf_desig", "").strip() or (meeting_details.get("company_name", "").strip() or "Client")
    p_conf_info = doc.add_paragraph(f"{conf_name}\n{conf_desig}")
    p_conf_info.paragraph_format.space_after = Pt(6)
    for r in p_conf_info.runs:
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(9.5)

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


def export_to_pdf_template_1(df: pd.DataFrame, meeting_details: dict, other_discussions: str) -> BytesIO:
    """Generate PDF for Template 1 (Standard Corporate Combined Table)."""
    if not HAS_REPORTLAB:
        return None
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    story, styles = [], getSampleStyleSheet()
    
    style_title = ParagraphStyle('DocTitle', parent=styles['Normal'], fontName=docfonts.FONT_HEAD_BOLD, fontSize=11.5, alignment=1, spaceAfter=2)
    company_target = meeting_details.get("external_attendees", [])
    primary_client_rep = company_target[0] if company_target else meeting_details.get("company_name", "").strip() or "CLIENT"
    style_subtitle = ParagraphStyle('DocSubTitle', parent=styles['Normal'], fontName=docfonts.FONT_HEAD_BOLD, fontSize=10.5, alignment=1, spaceAfter=10)
    style_body = ParagraphStyle('DocBody', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=9, leading=12, spaceAfter=3)
    style_th = ParagraphStyle('TableHead', parent=styles['Normal'], fontName=docfonts.FONT_HEAD_BOLD, fontSize=8.5, leading=10, alignment=1)
    style_td = ParagraphStyle('TableData', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=8, leading=10)
    style_td_center = ParagraphStyle('TableDataCenter', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=8, leading=10, alignment=1)

    story.append(Paragraph("<u>MINUTES OF THE MEETING</u>", style_title))
    story.append(Paragraph(f"PRIME PHILIPPINES & {primary_client_rep.upper()}", style_subtitle))
    
    date_str, time_str = meeting_details.get("date", "____________"), meeting_details.get("time_range", "")
    full_date = f"<b>Date:</b> {date_str}" + (f", {time_str}" if time_str.strip() else "")
    story.append(Paragraph(full_date, style_body))
    story.append(Paragraph(f"<b>Location:</b> {meeting_details.get('location', '____________')}", style_body))
    
    prime_atts, ext_atts = meeting_details.get("prime_attendees", []), meeting_details.get("external_attendees", [])
    att_list = []
    for att in ext_atts:
        if att.strip():
            att_list.append(f"{att}{f', {meeting_details.get('company_name')}' if meeting_details.get('company_name') else ''}")
    for att in prime_atts:
        att_list.append(f"{att} – PRIME Philippines")
    
    if att_list:
        story.append(Paragraph(f"<b>Attended by:</b>&nbsp;&nbsp;&nbsp;&nbsp;{att_list[0]}", style_body))
        for a in att_list[1:]:
            story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;{a}", style_body))
    
    story.append(Spacer(1, 4))
    client_display = meeting_details.get('company_name', '').strip() or "the Client"
    story.append(Paragraph(f"During the meeting held last {date_str}, PRIME Philippines, represented by the attendee/s shown above, met with {client_display} to discuss opportunities for collaboration.", style_body))
    story.append(Spacer(1, 6))
    
    table_data = [[Paragraph("<b>Discussion Points</b>", style_th), Paragraph("<b>Action Plan</b>", style_th), Paragraph("<b>Indicative Delivery Date</b>", style_th), Paragraph("<b>Person-in-charge</b>", style_th)]]
    for i, row in df.iterrows():
        table_data.append([
            Paragraph(f"{i+1}. {str(row.get('Discussion Points', ''))}", style_td),
            Paragraph(str(row.get("Action Plan", "")), style_td),
            Paragraph(str(row.get("Indicative Delivery Date", "")), style_td_center),
            Paragraph(str(row.get("Person-in-charge", "")), style_td_center)
        ])
    
    t = Table(table_data, colWidths=[2.4 * inch, 2.3 * inch, 1.1 * inch, 1.0 * inch], repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#C9AB4C')),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#69727d')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4)
    ]))
    story.append(t)
    
    note_style = ParagraphStyle('Note', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=7.5, leading=9, spaceBefore=4)
    story.append(Paragraph("*Note: The indicative delivery date serves as reference point and still subject to changes. Furthermore, it depends on the progress of both parties.", note_style))
    
    if other_discussions.strip():
        story.append(Spacer(1, 6))
        story.append(Paragraph("<b>Other Discussions:</b>", style_body))
        story.append(Paragraph(other_discussions, style_body))
    
    story.append(Spacer(1, 8))
    prep_name = meeting_details.get("prep_name", "").strip() or "____________________"
    prep_desig = meeting_details.get("prep_desig", "").strip() or "PRIME Philippines"
    conf_name = meeting_details.get("conf_name", "").strip() or (ext_atts[0] if ext_atts else "____________________")
    conf_desig = meeting_details.get("conf_desig", "").strip() or (meeting_details.get("company_name", "").strip() or "Client")
    
    sign_data = [
        [Paragraph("<b>Prepared by:</b>", style_body), Paragraph("<b>Confirmed by:</b>", style_body)],
        [Paragraph("_______________________________", style_body), Paragraph("_______________________________", style_body)],
        [Paragraph(f"{prep_name}<br/>{prep_desig}", style_body), Paragraph(f"{conf_name}<br/>{conf_desig}", style_body)]
    ]
    sign_table = Table(sign_data, colWidths=[3.4 * inch, 3.4 * inch])
    sign_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2)
    ]))
    story.append(sign_table)
    
    doc.build(story)
    buffer.seek(0)
    return buffer


# -------------------------------------------------------------
# TEMPLATE 2: Detailed General Meeting (Vertical Layout)
# -------------------------------------------------------------
def export_to_word_template_2(df: pd.DataFrame, meeting_details: dict, other_discussions: str) -> BytesIO:
    """Generate Word (.docx) for Template 2 (Detailed Sectional Report)."""
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("MINUTES OF THE MEETING")
    r_title.bold = True
    r_title.font.name = docfonts.WORD_BODY
    r_title.font.size = Pt(16)
    
    company_target = meeting_details.get("external_attendees", [])
    primary_client_rep = company_target[0] if company_target else meeting_details.get("company_name", "").strip() or "General Meeting"
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run(f"Project / Client: {primary_client_rep}")
    r_sub.font.name = docfonts.WORD_BODY
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(105, 114, 125)

    doc.add_heading('Meeting Details', level=2)
    details_table = doc.add_table(rows=5, cols=2)
    details_table.style = 'Table Grid'
    
    date_str = meeting_details.get("date", "____________")
    time_str = meeting_details.get("time_range", "____________")
    location_str = meeting_details.get('location', '____________')
    prep_name = meeting_details.get("prep_name", "").strip() or "____________________"
    
    details_map = [
        ("Date", date_str),
        ("Time", time_str),
        ("Venue", location_str),
        ("Prepared by", prep_name),
        ("Date prepared", datetime.datetime.now().strftime("%B %d, %Y"))
    ]
    
    for i, (key, val) in enumerate(details_map):
        cells = details_table.rows[i].cells
        cells[0].text = key
        cells[1].text = val
        cells[0].paragraphs[0].runs[0].font.bold = True
        for cell in cells:
            cell.paragraphs[0].runs[0].font.name = docfonts.WORD_BODY
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph()

    doc.add_heading('Attendees', level=2)
    all_atts = meeting_details.get("prime_attendees", []) + meeting_details.get("external_attendees", [])
    for att in all_atts:
        if att.strip():
            p_att = doc.add_paragraph(style='List Bullet')
            r_att = p_att.add_run(att.strip())
            r_att.font.name = docfonts.WORD_BODY
            r_att.font.size = Pt(10)

    doc.add_paragraph()

    doc.add_heading('Purpose & Summary', level=2)
    p_purp = doc.add_paragraph(other_discussions if other_discussions.strip() else "To discuss project updates, ongoing deliverables, and establish clear action plans.")
    p_purp.paragraph_format.space_after = Pt(12)
    for r in p_purp.runs:
        r.font.name, r.font.size = docfonts.WORD_BODY, Pt(10)

    doc.add_heading('Discussion Points', level=2)
    for i, row in df.iterrows():
        p_dp = doc.add_paragraph(style='List Number')
        r_dp = p_dp.add_run(str(row.get('Discussion Points', '')))
        r_dp.font.name = docfonts.WORD_BODY
        r_dp.font.size = Pt(10)

    doc.add_paragraph()

    doc.add_heading('Action Plan', level=2)
    act_table = doc.add_table(rows=len(df)+1, cols=4)
    act_table.style = 'Table Grid'
    act_headers = ["#", "Action Plan", "Owner", "Deadline"]
    col_widths = [Inches(0.5), Inches(3.5), Inches(1.5), Inches(1.5)]
    
    for i, header in enumerate(act_headers):
        cell = act_table.rows[0].cells[i]
        cell.width, cell.text = col_widths[i], header
        set_cell_shading(cell, "003366")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs: 
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.size, p.runs[0].font.name = Pt(9), docfonts.WORD_BODY

    for i, row in df.iterrows():
        cells = act_table.rows[i+1].cells
        cells[0].text = str(i+1)
        cells[1].text = str(row.get("Action Plan", ""))
        cells[2].text = str(row.get("Person-in-charge", ""))
        cells[3].text = str(row.get("Indicative Delivery Date", ""))
        for c_idx, cell in enumerate(cells):
            cell.width = col_widths[c_idx]
            p = cell.paragraphs[0]
            if c_idx in [0, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if p.runs:
                p.runs[0].font.size, p.runs[0].font.name = Pt(9), docfonts.WORD_BODY

    doc.add_paragraph()

    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.paragraph_format.space_before = Pt(24)
    r_footer = p_footer.add_run(f"Prepared for circulation to {primary_client_rep}. Please return corrections before this is treated as the agreed record.")
    r_footer.italic = True
    r_footer.font.color.rgb = RGBColor(105, 114, 125)
    r_footer.font.name = docfonts.WORD_BODY
    r_footer.font.size = Pt(8)

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


def export_to_pdf_template_2(df: pd.DataFrame, meeting_details: dict, other_discussions: str) -> BytesIO:
    """Generate PDF for Template 2 (Detailed Sectional Report)."""
    if not HAS_REPORTLAB:
        return None
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    story, styles = [], getSampleStyleSheet()
    
    style_title = ParagraphStyle('Title2', parent=styles['Normal'], fontName=docfonts.FONT_HEAD_BOLD, fontSize=15, alignment=1, spaceAfter=2)
    style_subtitle = ParagraphStyle('SubTitle2', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=10.5, textColor=colors.HexColor('#69727d'), alignment=1, spaceAfter=20)
    style_h2 = ParagraphStyle('Heading2', parent=styles['Normal'], fontName=docfonts.FONT_HEAD_BOLD, fontSize=12, textColor=colors.HexColor("#003366"), spaceBefore=12, spaceAfter=6)
    style_body = ParagraphStyle('Body2', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=9.5, leading=14, spaceAfter=4)
    style_th = ParagraphStyle('TH2', parent=styles['Normal'], fontName=docfonts.FONT_HEAD_BOLD, fontSize=9, textColor=colors.white, alignment=1)
    style_td = ParagraphStyle('TD2', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=9, leading=12)
    style_td_center = ParagraphStyle('TDC2', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=9, leading=12, alignment=1)
    
    company_target = meeting_details.get("external_attendees", [])
    primary_client_rep = company_target[0] if company_target else meeting_details.get("company_name", "").strip() or "General Meeting"

    story.append(Paragraph("MINUTES OF THE MEETING", style_title))
    story.append(Paragraph(f"Project / Client: {primary_client_rep}", style_subtitle))
    
    story.append(Paragraph("Meeting Details", style_h2))
    date_str = meeting_details.get("date", "____________")
    time_str = meeting_details.get("time_range", "____________")
    location_str = meeting_details.get('location', '____________')
    prep_name = meeting_details.get("prep_name", "").strip() or "____________________"
    
    details_data = [
        [Paragraph("<b>Date</b>", style_body), Paragraph(date_str, style_body)],
        [Paragraph("<b>Time</b>", style_body), Paragraph(time_str, style_body)],
        [Paragraph("<b>Venue</b>", style_body), Paragraph(location_str, style_body)],
        [Paragraph("<b>Prepared by</b>", style_body), Paragraph(prep_name, style_body)],
        [Paragraph("<b>Date prepared</b>", style_body), Paragraph(datetime.datetime.now().strftime("%B %d, %Y"), style_body)]
    ]
    t_details = Table(details_data, colWidths=[2 * inch, 4.5 * inch])
    t_details.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#69727d')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6)
    ]))
    story.append(t_details)
    story.append(Spacer(1, 10))

    story.append(Paragraph("Attendees", style_h2))
    all_atts = meeting_details.get("prime_attendees", []) + meeting_details.get("external_attendees", [])
    att_items = [ListItem(Paragraph(a.strip(), style_body)) for a in all_atts if a.strip()]
    if att_items:
        story.append(ListFlowable(att_items, bulletType='bullet'))
    story.append(Spacer(1, 10))

    story.append(Paragraph("Purpose & Summary", style_h2))
    story.append(Paragraph(other_discussions if other_discussions.strip() else "To discuss project updates, ongoing deliverables, and establish clear action plans.", style_body))
    story.append(Spacer(1, 10))

    story.append(Paragraph("Discussion Points", style_h2))
    dp_items = [ListItem(Paragraph(str(row.get('Discussion Points', '')), style_body)) for _, row in df.iterrows()]
    if dp_items:
        story.append(ListFlowable(dp_items, bulletType='1'))
    story.append(Spacer(1, 10))

    story.append(Paragraph("Action Plan", style_h2))
    act_data = [[Paragraph("<b>#</b>", style_th), Paragraph("<b>Action Plan</b>", style_th), Paragraph("<b>Owner</b>", style_th), Paragraph("<b>Deadline</b>", style_th)]]
    for i, row in df.iterrows():
        act_data.append([
            Paragraph(str(i+1), style_td_center),
            Paragraph(str(row.get("Action Plan", "")), style_td),
            Paragraph(str(row.get("Person-in-charge", "")), style_td_center),
            Paragraph(str(row.get("Indicative Delivery Date", "")), style_td_center)
        ])
    
    t_act = Table(act_data, colWidths=[0.4 * inch, 3.5 * inch, 1.3 * inch, 1.3 * inch], repeatRows=1)
    t_act.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#003366')),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#69727d')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5)
    ]))
    story.append(t_act)

    story.append(Spacer(1, 24))
    footer_style = ParagraphStyle('Footer2', parent=styles['Normal'], fontName=docfonts.FONT_BODY, fontSize=8, textColor=colors.HexColor('#69727d'), alignment=1)
    story.append(Paragraph(f"Prepared for circulation to {primary_client_rep}. Please return corrections before this is treated as the agreed record.", footer_style))

    doc.build(story)
    buffer.seek(0)
    return buffer
